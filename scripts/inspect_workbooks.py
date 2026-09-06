"""Read-only domain evidence extraction; excludes names and personal details."""

import hashlib, json
from pathlib import Path
import openpyxl

out = []
for name in ["OK 505SRG.xlsx", "Strength 2.xlsx"]:
    path = Path.home() / "Downloads" / name
    book = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    item = {
        "file": name,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "external_links": len(book._external_links),
        "sheets": [],
    }
    for sheet in book:
        formulas = [
            (c.coordinate, c.value) for row in sheet for c in row if c.data_type == "f"
        ]
        errors = [
            (c.coordinate, c.value)
            for row in cached[sheet.title]
            for c in row
            if c.data_type == "e"
        ]
        item["sheets"].append(
            {
                "name": sheet.title,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "formula_count": len(formulas),
                "formula_examples": formulas[:4],
                "cached_error_count": len(errors),
                "error_examples": errors[:4],
            }
        )
    out.append(item)
Path(".local").mkdir(exist_ok=True)
Path(".local/workbook-evidence.json").write_text(json.dumps(out, indent=2))
print(json.dumps(out, indent=2))
